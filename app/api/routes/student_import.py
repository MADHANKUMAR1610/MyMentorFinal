from uuid import UUID
from io import BytesIO

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    UploadFile,
    status,
)



from openpyxl import load_workbook

from sqlalchemy.ext.asyncio import AsyncSession

from app.api.dependencies import (
    get_current_system_admin,
)

from app.database.database import get_db

from app.models.college import College
from app.models.user import User

from app.repositories.college_repository import (
    CollegeRepository,
)

from app.repositories.user_repository import (
    UserRepository,
)
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from io import BytesIO
from fastapi import APIRouter

router = APIRouter(
    prefix="/students/import",
    tags=["Student Import"],
)
@router.post(
    "",
)
async def import_students_excel(
    college_id: UUID,
    file: UploadFile = File(...),
    current_user: User = Depends(
        get_current_system_admin
    ),
    session: AsyncSession = Depends(
        get_db
    ),
):

    # --------------------------------------------------------
    # VALIDATE FILE
    # --------------------------------------------------------

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="Excel file is required.",
        )

    if not file.filename.lower().endswith(
        (".xlsx", ".xlsm")
    ):
        raise HTTPException(
            status_code=400,
            detail="Please upload an Excel file.",
        )

    # --------------------------------------------------------
    # GET COLLEGE
    # --------------------------------------------------------

    college_repository = CollegeRepository(
        session
    )

    college = await college_repository.get_by_id(
        college_id
    )

    if college is None:
        raise HTTPException(
            status_code=404,
            detail="College not found.",
        )

    # --------------------------------------------------------
    # READ EXCEL
    # --------------------------------------------------------

    content = await file.read()

    try:
        workbook = load_workbook(
            BytesIO(content),
            data_only=True,
        )

    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Unable to read the Excel file.",
        )

    worksheet = workbook.active

    # --------------------------------------------------------
    # VALIDATE HEADERS
    # --------------------------------------------------------

    expected_headers = [
        "Name",
        "Email",
        "Phone",
        "Department",
        "Year",
    ]

    actual_headers = [
        str(cell.value).strip()
        if cell.value is not None
        else ""
        for cell in worksheet[1]
    ]

    if actual_headers != expected_headers:

        raise HTTPException(
            status_code=400,
            detail={
                "message": "Invalid Excel columns.",
                "expected_columns": expected_headers,
                "received_columns": actual_headers,
            },
        )

    # --------------------------------------------------------
    # USER REPOSITORY
    # --------------------------------------------------------

    user_repository = UserRepository(
        session
    )

    imported_students = []
    errors = []

    # --------------------------------------------------------
    # PROCESS ROWS
    # --------------------------------------------------------

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):

        # Skip completely empty rows
        if not any(
            value is not None
            for value in row
        ):
            continue

        name = (
            str(row[0]).strip()
            if row[0] is not None
            else ""
        )

        email = (
            str(row[1]).strip().lower()
            if row[1] is not None
            else ""
        )

        phone = (
            str(row[2]).strip()
            if row[2] is not None
            else ""
        )

        department = (
            str(row[3]).strip()
            if row[3] is not None
            else ""
        )

        year = (
            str(row[4]).strip()
            if row[4] is not None
            else ""
        )

        # ----------------------------------------------------
        # REQUIRED FIELD VALIDATION
        # ----------------------------------------------------

        if not name:
            errors.append(
                {
                    "row": row_number,
                    "error": "Name is required.",
                }
            )
            continue

        if not email:
            errors.append(
                {
                    "row": row_number,
                    "error": "Email is required.",
                }
            )
            continue

        # ----------------------------------------------------
        # DUPLICATE EMAIL
        # ----------------------------------------------------

        existing_user = (
            await user_repository.get_by_email(
                email
            )
        )

        if existing_user is not None:

            errors.append(
                {
                    "row": row_number,
                    "email": email,
                    "error": (
                        "A user with this email "
                        "already exists."
                    ),
                }
            )

            continue

        # ----------------------------------------------------
        # CREATE STUDENT
        # ----------------------------------------------------

        student = User(
            name=name,
            email=email,
            phone=phone or None,
            department=department or None,
            year=year or None,
            role="student",
            college_id=college.id,
            is_active=True,
            is_verified=True,
            onboarded=False,
        )

        created_student = await (
            user_repository.create(
                student
            )
        )

        imported_students.append(
            {
                "id": str(created_student.id),
                "name": created_student.name,
                "email": created_student.email,
                "phone": created_student.phone,
                "department": created_student.department,
                "year": created_student.year,
                "college_id": str(college.id),
                "student_code": None,
            }
        )

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    await session.commit()

    return {
        "message": "Student Excel imported successfully.",
        "college": {
            "id": str(college.id),
            "name": college.name,
            "code": college.code,
        },
        "imported_count": len(
            imported_students
        ),
        "students": imported_students,
        "errors": errors,
    }
@router.post(
    "/{college_id}/generate-codes",
)
async def generate_student_codes(
    college_id: UUID,
    current_user: User = Depends(
        get_current_system_admin
    ),
    session: AsyncSession = Depends(
        get_db
    ),
):

    college_repository = CollegeRepository(
        session
    )

    college = await college_repository.get_by_id(
        college_id
    )

    if college is None:
        raise HTTPException(
            status_code=404,
            detail="College not found.",
        )

    user_repository = UserRepository(
        session
    )

    students = await (
        user_repository.get_students_by_college(
            college_id
        )
    )

    # --------------------------------------------------------
    # EXISTING GENERATED CODES
    # --------------------------------------------------------

    students_without_code = [
        student
        for student in students
        if not student.student_code
    ]

    if not students_without_code:

        return {
            "message": (
                "All students already have "
                "student codes."
            ),
            "college": {
                "id": str(college.id),
                "name": college.name,
                "code": college.code,
            },
            "generated_count": 0,
            "students": [
                {
                    "id": str(student.id),
                    "name": student.name,
                    "email": student.email,
                    "phone": student.phone,
                    "department": student.department,
                    "year": student.year,
                    "student_code": (
                        student.student_code
                    ),
                }
                for student in students
            ],
        }

    # --------------------------------------------------------
    # FIND NEXT NUMBER
    # --------------------------------------------------------

    existing_numbers = []

    prefix = f"{college.code}-"

    for student in students:

        if (
            student.student_code
            and student.student_code.startswith(prefix)
        ):

            number_part = (
                student.student_code[
                    len(prefix):
                ]
            )

            if number_part.isdigit():

                existing_numbers.append(
                    int(number_part)
                )

    next_number = (
        max(existing_numbers) + 1
        if existing_numbers
        else 1
    )

    # --------------------------------------------------------
    # GENERATE
    # --------------------------------------------------------

    generated = []

    for student in students_without_code:

        student.student_code = (
            f"{college.code}-{next_number:06d}"
        )

        generated.append(student)

        next_number += 1

    await session.commit()

    # --------------------------------------------------------
    # RESPONSE
    # --------------------------------------------------------

    return {
        "message": (
            "Student codes generated successfully."
        ),
        "college": {
            "id": str(college.id),
            "name": college.name,
            "code": college.code,
        },
        "generated_count": len(
            generated
        ),
        "students": [
            {
                "id": str(student.id),
                "name": student.name,
                "email": student.email,
                "phone": student.phone,
                "department": student.department,
                "year": student.year,
                "student_code": (
                    student.student_code
                ),
            }
            for student in students
        ],
    }
@router.get(
    "/{college_id}/export",
)
async def export_students_excel(
    college_id: UUID,
    current_user: User = Depends(
        get_current_system_admin
    ),
    session: AsyncSession = Depends(
        get_db
    ),
):

    college_repository = CollegeRepository(
        session
    )

    college = await college_repository.get_by_id(
        college_id
    )

    if college is None:
        raise HTTPException(
            status_code=404,
            detail="College not found.",
        )

    user_repository = UserRepository(
        session
    )

    students = await (
        user_repository.get_students_by_college(
            college_id
        )
    )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Students"

    headers = [
        "Name",
        "Email",
        "Phone",
        "Department",
        "Year",
        "Student Code",
    ]

    for column_index, header in enumerate(
        headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=header,
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    for row_index, student in enumerate(
        students,
        start=2,
    ):

        worksheet.cell(
            row=row_index,
            column=1,
            value=student.name,
        )

        worksheet.cell(
            row=row_index,
            column=2,
            value=student.email,
        )

        worksheet.cell(
            row=row_index,
            column=3,
            value=student.phone,
        )

        worksheet.cell(
            row=row_index,
            column=4,
            value=student.department,
        )

        worksheet.cell(
            row=row_index,
            column=5,
            value=student.year,
        )

        worksheet.cell(
            row=row_index,
            column=6,
            value=student.student_code,
        )

    worksheet.column_dimensions["A"].width = 30
    worksheet.column_dimensions["B"].width = 35
    worksheet.column_dimensions["C"].width = 18
    worksheet.column_dimensions["D"].width = 25
    worksheet.column_dimensions["E"].width = 15
    worksheet.column_dimensions["F"].width = 25

    worksheet.freeze_panes = "A2"

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; '
                f'filename="students_{college.code}.xlsx"'
            )
        },
    )